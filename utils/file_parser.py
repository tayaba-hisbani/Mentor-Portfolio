from io import BytesIO

from pypdf import PdfReader
from docx import Document


def extract_text(uploaded_file) -> str:
    name = uploaded_file.name.lower()
    raw = uploaded_file.getvalue()

    try:
        if name.endswith(".pdf"):
            return _extract_pdf(raw)
        elif name.endswith(".docx"):
            return _extract_docx(raw)
        elif name.endswith((".txt", ".md")):
            return raw.decode("utf-8", errors="ignore")
        elif name.endswith((".png", ".jpg", ".jpeg", ".webp")):
            return (
                f"[Image file '{uploaded_file.name}' uploaded. Text content "
                "could not be extracted automatically — describe what this "
                "image shows in the notes box if you'd like it reviewed.]"
            )
        else:
            return (
                f"[Unsupported file type for '{uploaded_file.name}'. "
                "Please upload a PDF, DOCX, TXT, or MD file.]"
            )
    except Exception as exc:
        return f"[Could not read '{uploaded_file.name}': {exc}]"


def _extract_pdf(raw_bytes: bytes) -> str:
    reader = PdfReader(BytesIO(raw_bytes))
    pages = [page.extract_text() or "" for page in reader.pages]
    text = "\n".join(pages).strip()
    return text or "[No extractable text found in this PDF — it may be scanned images.]"


def _extract_docx(raw_bytes: bytes) -> str:
    doc = Document(BytesIO(raw_bytes))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    text = "\n".join(parts).strip()
    return text or "[No extractable text found in this DOCX.]"
